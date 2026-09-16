from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pathlib import Path

def add_text(slide,text,x,y,w,h,size=24,bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=str(text); p.font.size=Pt(size); p.font.bold=bold
    return box

def build_ppt(slides,output_path,title,theme="Professional"):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    for i,s in enumerate(slides):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=RGBColor(248,250,252)
        add_text(slide,s.get("title","Untitled"),.65,.35,12,.7,28,True)
        bullets=s.get("bullets",[])
        body="\n".join("• "+str(x) for x in bullets)
        add_text(slide,body,.8,1.35,7.2,5.2,19)
        img=s.get("image_url") or s.get("image_path")
        if img:
            p=Path(img)
            if p.exists(): slide.shapes.add_picture(str(p),Inches(8.35),Inches(1.45),width=Inches(4.35))
        notes=s.get("speaker_notes","")
        if notes:
            add_text(slide,"Speaker notes:\n"+notes,8.35,6.25,4.35,.75,9,False)
    prs.save(output_path)
